"""Jadvallarni chiroyli, rangli Excel (.xlsx) ga eksport qilish (openpyxl)."""

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Tag/signal ranglari (fon, matn) — saytdagi ranglarga mos
RANG = {
    'green': ('DCFCE7', '166534'),
    'red': ('FEE2E2', '991B1B'),
    'amber': ('FEF3C7', '92400E'),
    'orange': ('FFEDD5', '9A3412'),
}
_MONEY = re.compile(r'^\$\s*([\d,]+\.\d{1,2})$')


def qil_xlsx(sarlavha, headers, rows):
    """headers: [str]; rows: [[(matn, rang|None), ...], ...] -> xlsx bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = (sarlavha or 'Hisobot')[:31]

    head_fill = PatternFill('solid', fgColor='047857')
    head_font = Font(bold=True, color='FFFFFF', size=11)
    thin = Side(style='thin', color='E5E9F0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center')

    widths = [len(str(h)) for h in headers]

    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=ci, value=str(h))
        c.fill = head_fill
        c.font = head_font
        c.border = border
        c.alignment = center

    for ri, row in enumerate(rows, start=2):
        for ci, (matn, rang) in enumerate(row, start=1):
            matn = matn or ''
            m = _MONEY.match(matn)
            if m:
                cell = ws.cell(row=ri, column=ci, value=float(m.group(1).replace(',', '')))
                cell.number_format = '"$"#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell = ws.cell(row=ri, column=ci, value=matn)
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
            cell.border = border
            if rang and rang in RANG:
                fg, fc = RANG[rang]
                cell.fill = PatternFill('solid', fgColor=fg)
                cell.font = Font(color=fc, bold=True)
            elif ri % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F7FAF9')
            if len(matn) > widths[ci - 1]:
                widths[ci - 1] = len(matn)

    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(max(w + 3, 11), 48)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
